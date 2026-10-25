from pathlib import Path
from shutil import copy2
import json,re
from openpyxl import load_workbook
from docx import Document

R=Path(r'D:\git\RZS_Repository'); W=Path(__file__).parent;changes=[]
def backup(p):
 dst=W/'原文备份'/p.relative_to(R);dst.parent.mkdir(parents=True,exist_ok=True)
 if not dst.exists():copy2(p,dst)
def save_doc(path,replace_map,banner):
 p=R/path;backup(p);d=Document(p);count=0
 def update(para):
  nonlocal count
  old=para.text;new=old
  for a,b in replace_map.items():new=new.replace(a,b)
  if old!=new:
   # 保留段落中的图片和对象，纯文字变更优先在同一run内完成。
   changed=False
   for run in para.runs:
    before=run.text;after=before
    for a,b in replace_map.items():after=after.replace(a,b)
    if before!=after:run.text=after;changed=True
   if para.text!=new and not para._p.xpath('.//w:drawing | .//w:pict | .//w:object'):
    if para.runs:
     para.runs[0].text=new
     for run in para.runs[1:]:run.text=''
    else:para.add_run(new)
   count+=1
 for para in list(d.paragraphs):update(para)
 for t in d.tables:
  for row in t.rows:
   for c in row.cells:
    for para in c.paragraphs:update(para)
 if banner and not any(banner in p.text for p in d.paragraphs):
  first=d.paragraphs[0];first.insert_paragraph_before(banner)
 d.save(p);changes.append({'file':path,'updated_texts':count,'note':'原文备份已保存；图片对象保留'})

path='低压盒/硬件电路接口.xlsx';p=R/path;backup(p);w=load_workbook(p)
s=w['接口表']
s['E15']='板端5 V逻辑高有效；外部VCU 12 V命令先由基板CN11.3经R7/D21/R41限流钳位至约4.7 V，再到基板H4.2和本板H1.2。禁止把12 V直接接入本板逻辑脚。故障/开路及掉电默认态须实测。'
s['E16']='安全回路经外部DCDC后为12 V。掉电使RLY2释放；RLY2为SRD-12VDC-SL-C。EBS_PWR/EBS_WORK主电源支路仍为24 V，RLY1为24 V。'
s['E17']='EBS_PWR为24 V，来自最终整车图定义的RES/ASMS后主电源支路。主电源丢失应撤销EBS电磁阀能量许可；RLY1为24 V，RLY2为12 V。机械制动、保持和释放需气路验收。'
s['E22']='正常约12 V、故障0 V；由VCU独立输出，经基板CN11.2/H9.2送本板H2.2，20 kΩ/12 kΩ分压。TSMS上电前须先由有效健康状态建立12 V。删除双路健康后才开始监测的等待设计，实际R_E故障锁存保留。不是车载PB13的3.3 V FAULT_OUT。'
s['I22']='5 V高边状态输出，关闭时可能悬空。基板H17.3 R_E_LED按设计预留未使用；EBS/RES错误灯由基板CN7.16控制，不能依赖本输出已接入错误灯。'
for row in range(29,49):
 for cell in s[row]:
  if isinstance(cell.value,str):cell.value=cell.value.replace('U11.','U14.').replace('CN1.','U15.')
for row in range(5,14):
 for cell in s[row]:
  if isinstance(cell.value,str):cell.value=cell.value.replace('故障/开路：低于目标阈值为电流条件无效','故障/开路：板内10 kΩ上拉，开路趋向高电平；是否断开取决于油压与电流组合及延时，须故障注入').replace('故障/开路：低于目标阈值为制动条件无效','故障/开路：板内10 kΩ上拉，开路趋向高电平；是否断开取决于电流组合及延时，须故障注入')
c=w['校核清单']
c['C5']='VCU_WDI为0～5 V推挽10 Hz；EBS_TRIG连接器侧12 V，经基板调理后进入EBS 5 V逻辑域'
c['D5']='CN11.3→R7/D21 4.7 V钳位→R41→基板H4.2；WDI直接转接；反馈幅值与极性待测'
c['E5']='分别测命令侧12 V与板端约4.7 V、WDI停止/错频动作、EBS_CHECK/VCU_FB完整真值表'
c['C8']='旧Q2源漏疑问已由当前网表澄清；实物需按版号核对'
c['D8']='当前网表Q2.2 Source接GND；Q2.3 Drain接C3/R10/U7定时节点'
c['E8']='按AO3400实际封装、PCB和焊接核对；不再引用废弃反接描述';c['F8']='网表已澄清，实物待核对'
c['C15']='RESET为低有效开漏接口；板端5 V上拉，基板另有约4.7 V共用偏置'
row=c.max_row+1
for vals in [('高','EBS','SDC_IN按12 V新方案','外部DCDC→12 V RLY2；主支路24 V RLY1','验证DCDC、线束、触点及失电制动','待实物确认'),('高','TSAL','传感器量程与板内参考未一致','基板0～500 V→1～5 V时60 V对应1.48 V，板内参考0.300/0.743 V','冻结完整传递函数和隔离/调理，扫描实际翻转点','待系统确认')]:
 for j,v in enumerate(vals,1):c.cell(row,j,v)
 row+=1
w.save(p);changes.append({'file':path,'note':'按最新EBS/RES方案、现relay/bspd网表修正；保留待测状态'})

path='基板/基板全部接口特性与定义_2026-07-31_v7.xlsx';p=R/path;backup(p);w=load_workbook(p)
for sheet in w:
 for row in sheet:
  for cell in row:
   if not isinstance(cell.value,str):continue
   v=cell.value
   v=v.replace('安全回路名义24V','经DCDC后的12V安全回路').replace('名义24V安全回路','经DCDC后的12V安全回路')
   v=v.replace('开路/低于阈值通常判制动条件无效，具体故障逻辑由BSPD板决定','低于阈值为制动条件无效；开路由BSPD板内10k上拉趋向有效高，需结合电流与延时验证')
   v=v.replace('低于阈值表示电流条件无效；开路判据由BSPD板决定','低于阈值表示电流条件无效；开路由BSPD板内10k上拉趋向高，需结合油压与延时验证')
   if v=='外部电平直通，低有效故障':v='健康12V／故障0V；外部直通'
   if v=='外部电平直通至锁存板':v='健康12V／故障0V；由VCU独立输出'
   if v=='1~5V隔离模拟量':v='原表1～5V，比例与TSAL阈值待核定'
   if v=='当前系统定义0~500V→1~5V':v='源表0～500V→1～5V；与TSAL阈值冲突，待核定'
   cell.value=v
s=w['总览'];s['A1']='基板全部接口特性与定义 2026-10-16校订';s['A2']='连接依据为网表/base.tel；功能与电压采用2026-10-04新方案。SDC_IN_EBS为外部DCDC后12V，EBS主支路24V；RES_ERROR健康12V、故障0V且先于TSMS。实物验证状态另列。'
s=w['外部接口CN7-CN12']
for row in s:
 if row[0].value=='CN11' and str(row[1].value)=='2':row[6].value='VCU健康输出12V、故障0V；TSMS前先建立有效健康态；直通H9.2，由锁存板调理'
 if row[0].value=='CN11' and str(row[1].value)=='10':row[11].value='新方案12V；DCDC与实物待验收'
s=w['板卡排针H1-H20']
for row in s:
 if row[0].value=='H9' and str(row[1].value)=='1':row[4].value='输入（相对锁存板）';row[6].value='来自基板H3.3 EBS_ERROR'
s=w['待确认与测试'];r=s.max_row+1
for vals in [('P12','高','TSAL传感器比例与比较阈值冲突','原表0～500V→1～5V；60V为1.48V，TSAL参考0.300/0.743V','不能据此宣称60V判定通过','冻结完整传递函数并实测；不擅自更改量程','待系统确认'),('P13','高','EBS安全回路改为12V','2026-10-04方案及ebs.tel RLY2 12V','不得仍接旧24V','确认DCDC位置、SDC_IN与两继电器触点','待实物确认'),('P14','高','VCU输出高低边与基板适配','VCU表OUT4低边而CN7.16需12V高有效；resl_err也列低边','接线或软件极性可能不匹配','联合核对实际驱动与故障默认态','待系统确认')]:
 for j,v in enumerate(vals,1):s.cell(r,j,v)
 r+=1
w.save(p);changes.append({'file':path,'note':'新方案电压、RES来源、BSPD开路、H9方向及TSAL冲突校订'})

save_doc('继电器检测/继电器检测说明.docx',{'U11':'U14','CN1':'U15'},'2026年10月16日校订  连接器位号按当前网表/relay.tel更新为U14输入与供电、U15逻辑输出。旧版U11/CN1不再作为现用接线依据；实板仍须核对版号和丝印。')
save_doc('EBS/EBS功能说明（2026）.docx',{
'EBS触发：EBS_TRIG（VCU 12 V高有效，经12 V→5 V输入调理）':'EBS触发：外部VCU 12 V高有效命令先经基板R7/D21/R41调理，本板EBS_TRIG为约4.7 V高有效逻辑',
'VCU 12 V高有效触发；须经12 V→5 V调理后进入施密特反相器':'板端5 V逻辑高有效；外部12 V命令须先经基板R7/D21/R41调理至约4.7 V',
'安全回路电压；名义24 V，掉电时RLY2释放':'安全回路经外部DCDC后12 V，驱动12 V线圈RLY2；掉电释放',
'板子上电时解除刹车':'电气上电仅建立释放许可；已激活的EBS机械保持须按人工释放流程解除，不允许上电自动释放',
'EBS复位时解除刹车':'电气复位仅在故障解除后恢复许可，不代替EBS机械释放或AS Ready前检查',
'板子电压异常时接通外部电容，通过DCDC后驱动继电器刹车':'备用电容供电与欠压动作属于待核对系统需求，当前网表及实物闭环未确认，不能列为已实现功能',
'气压异常时刹车':'气压异常制动由整车气路/控制实现，板卡独立检测路径尚需核对，不列为本板已验证功能',
'EBS_PWR 旧文档写 12 V':'EBS_PWR为24 V，SDC_IN为12 V',
},'2026年10月16日校订  以2026年10月4日EBS说明和当前网表为准：EBS_PWR/EBS_WORK及RLY1主支路24 V，SDC_IN经外部DCDC为12 V、RLY2为12 V。删除EBS/RES双健康等待监测，实际故障锁存保留；TSMS前建立VCU RES_ERROR健康12 V。以下旧规则摘录仅供历史背景，不构成2026规则逐条审查或整车验收结论。')
save_doc('与其他组/硬件电路和整车控制有交集的地方.docx',{
'位于控制箱内，12V供电，所需信号为电机电流信号、2个油压传感器信号':'位于控制箱内，BSPD为24V供电，所需信号为电机电流信号和4路油压传感器信号',
'故障时为低电平':'IMD健康约12V/故障约9V，BMS健康约12V/故障约10V；EBS健康约5V/故障0V，RES健康约12V/故障0V',
'IMD/BMS锁存：':'IMD BMS EBS RES错误锁存',
'需与电池箱内电压传感器、高压盒内电压传感器共地':'仅传感器隔离后的低压输出侧与TSAL低压地相连；禁止将牵引高压侧接到低压地',
'使安全回路500ms后保持断开':'使安全回路在持续冲突后500ms内断开并保持',
},'2026年10月16日校订  本文现用接口按BSPD定义、TSAL定义、错误锁存说明及EBS说明更新。BSPD采用24V、四路油压；TSAL传感器到60V的比例尚待核定；旧电检步骤只作历史背景，禁止据此直接开展高压操作或认定2026合规。')
save_doc('VCU/方程式/方程式/FSAC测试台架说明.docx',{
'RES（远程制动系统）':'RES（远程急停系统）',
'（高压继电器线圈通电断开）':'（撤销安全回路许可使高压继电器释放断开，需按最终线圈/触点图核对）',
'OP1/2_IN：油压传感器信号':'OIL1～OIL4：四路油压传感器信号',
'高压继电器断路标志信号':'继电器状态一致性或断线信号，旧名与现ACTUAL/#需对照',
},'2026年10月16日校订  台架按现用网表及新接口方案：BSPD 24V四路油压；EBS主支路24V、安全回路DCDC后12V；RES G0B1/Si4463及CAN 0x510/0x511/0x512。图中旧接口名仅供追溯，完整逐针定义见设计报告/答辩准备_20261016/全车接口与信号索引.md。')
save_doc('设计报告/设计报告.docx',{},'2026年10月16日资料入口更新  本文件为早期概述，当前答辩和接口定义统一参见同目录“答辩准备_20261016/RZS整车电气系统答辩准备材料.docx”及“全车接口与信号索引.md”。旧图和旧功能概括不能代替现版接线与验收。')
(W/'legacy_corrections.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(changes,ensure_ascii=False,indent=2))
