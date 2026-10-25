from pathlib import Path
from zipfile import ZipFile
from openpyxl import load_workbook
from docx import Document
import json,hashlib,re
R=Path(r'D:\git\RZS_Repository');O=Path(__file__).parent.parent;W=O/'_work'
manifest=json.loads((W/'source_manifest.json').read_text(encoding='utf8'))
assert len(manifest)==22
assert all(not x['path'].startswith('VCU/') and 'VCU_INTEGRATION' not in x['path'] for x in manifest)
assert all(hashlib.sha256((R/x['path']).read_bytes()).hexdigest()==x['sha256'] for x in manifest)
d=Document(O/'RZS整车电气系统答辩准备材料.docx')
content='\n'.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
assert 'NXPS32K' not in content and '0x1E4' not in content and '150ms' not in content
assert 'VCU接口映射' not in content and 'CAN契约与VCU' not in content
assert all(s in content for s in ['ROTATE_OUT','ROTATE_IN','ROTATE_VCU','0.755','1.116','SDC_IN','R_E_RESET','高频问答','可维护性'])
assert all(not re.search(r'\[S(\d+)\]',x) or all(1<=int(n)<=22 for n in re.findall(r'\[S(\d+)\]',x)) for x in content.splitlines())
doc_media=[]
for original in (W/'原文备份').rglob('*.docx'):
 current=R/original.relative_to(W/'原文备份')
 with ZipFile(original) as a,ZipFile(current) as b:
  media=lambda z:{p:hashlib.sha256(z.read(p)).hexdigest() for p in z.namelist() if p.startswith('word/media/') and not p.endswith('/')}
  assert media(a)==media(b),str(current)
  doc_media.append({'file':str(current.relative_to(R)),'images_preserved':len(media(a))})
b=load_workbook(R/'基板/基板全部接口特性与定义_2026-07-31_v7.xlsx',data_only=True)
s=b['外部接口CN7-CN12']
pin_rows={str(row[1].value):row for row in s if row[0].value=='CN7'}
assert [pin_rows[p][2].value for p in ['11','13','14']]==['ROTATE_OUT','ROTATE_IN','ROTATE_VCU']
assert '不是GND' in s['A2'].value
a=load_workbook(R/'低压盒/硬件电路接口.xlsx',data_only=True)
assert '12 V' in a['接口表']['E16'].value and '24 V' in a['接口表']['E16'].value
assert '正常约12 V' in a['接口表']['E22'].value
audit=json.loads((W/'connection_audit.json').read_text(encoding='utf8'))
assert audit['external']==72 and audit['headers']==47
pdfpages=json.loads((W/'page_metrics.json').read_text(encoding='utf8'))
main_pages=[p for p in pdfpages if p['file']=='RZS整车电气系统答辩准备材料']
assert all(p['chars']>200 for p in main_pages),[p for p in main_pages if p['chars']<=200]
report={'sources':22,'main_pages':len(main_pages),'external_pins':72,'board_pins':47,'source_hashes_match':True,'VCU_program_excluded':True,'original_images':doc_media,'new_scheme_electrical_tests':'not performed'}
(W/'final_verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(report,ensure_ascii=False,indent=2))
