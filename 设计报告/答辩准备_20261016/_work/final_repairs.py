from pathlib import Path
from copy import copy
from openpyxl import load_workbook
from docx import Document
R=Path(r'D:\git\RZS_Repository');W=Path(__file__).parent
p=R/'基板/基板全部接口特性与定义_2026-07-31_v7.xlsx';w=load_workbook(p)
w['外部接口CN7-CN12']['A2']='连接关系以网表/base.tel（2026-10-04）为准。CN7.11/13/14为ROTATE_OUT/IN/VCU接K15，不是GND；CN8.6仍为GND。G5Q 2为COM、3为NO、4为NC；K15引脚与额定需按实际器件核对。'
for name in ['外部接口CN7-CN12','板卡排针H1-H20']:
 s=w[name]
 for row in s:
  if len(row)>2 and row[2].value=='GND':
   idx=10 if name.startswith('外部') else 7
   row[idx].value='GND 公共低压地；完整节点见网表/base.tel'
   s.row_dimensions[row[0].row].height=48
s=w['待确认与测试']
for row in list(s):
 if row[0].value=='P14':s.delete_rows(row[0].row,1)
for row in range(16,s.max_row+1):
 for col in range(1,8):s.cell(row,col)._style=copy(s.cell(14,col)._style)
w.save(p)
p=R/'低压盒/硬件电路接口.xlsx';w=load_workbook(p);s=w['校核清单']
for row in range(16,s.max_row+1):
 for col in range(1,7):s.cell(row,col)._style=copy(s.cell(14,col)._style)
w.save(p)

# 修正旧EBS文档残留的未引出SDC_OUT说法与调理位置。
p=R/'EBS/EBS功能说明（2026）.docx';d=Document(p)
repls={
 '安全回路输出:SDC_OUT':'安全回路输入SDC_IN用于驱动RLY2；现网表未引出名为SDC_OUT的独立连接器输出',
 'EBS_TRIG为12 V高有效，须经12 V→5 V输入调理':'外部EBS_TRIG命令为12 V高有效，由基板R7/D21/R41调理至约4.7 V后进入本板5 V逻辑域',
 '须经12 V→5 V输入调理':'外部12 V命令先经基板调理至约4.7 V，再送本板5 V逻辑域',
}
for para in d.paragraphs:
 old=para.text;new=old
 for a,b in repls.items():new=new.replace(a,b)
 if old!=new:
  para.runs[0].text=new
  for run in para.runs[1:]:run.text=''
d.save(p)

# 控制20行接口摘要的长度，防止引文或结尾说明孤页。
p=W/'build_material.py';s=p.read_text(encoding='utf8')
s=s.replace("electrical if len(electrical)<=38 else electrical[:36]+'…'", "electrical if len(electrical)<=22 else electrical[:20]+'…'")
s=s.replace("beh if len(beh)<=54 else beh[:52]+'…'", "beh if len(beh)<=34 else beh[:32]+'…'")
s=s.replace("table(['针脚','信号','电平或类型','作用与有效条件'],rows,[1.5,4.1,4.0,7.8],8.1)", "table(['针脚','信号','电平或类型','作用与有效条件'],rows,[1.5,4.1,4.0,7.8],8.5)")
s=s.replace("('另行核对的实现文件')", "('另行核对的实现文件')")
p.write_text(s,encoding='utf8')
print('Completed final source and pagination repairs')
