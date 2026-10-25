from pathlib import Path
import json,hashlib
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

root=Path(r'D:\git\RZS_Repository')
out=Path(__file__).parent
files=[
'基板/基板全部接口特性与定义_2026-07-31_v7.xlsx','低压盒/硬件电路接口.xlsx',
'VCU/方程式/方程式/报文id(1)(1).xlsx','VCU/方程式/方程式/VCU部分引脚改(1).xlsx','VCU/方程式/方程式/NXPS32K_引脚分配(1).xlsx',
'继电器检测/继电器检测说明.docx','EBS/EBS功能说明（2026）.docx','与其他组/硬件电路和整车控制有交集的地方.docx','与其他组/无人系统驾驶表.docx','设计报告/设计报告.docx','VCU/方程式/方程式/FSAC测试台架说明.docx',
'所有板子完整仿真/simulations/simulation_report.md','所有板子完整仿真/仿真/STATUS.md','ARDUINO测试/测试汇总.md',
'BSPD/BSPD定义.md','TSAL/定义.md','错误锁存/说明.txt','EBS/EBS说明.md',
'RES/firmware/docs/RES_SYSTEM_DESCRIPTION_CN.md','RES/firmware/docs/VCU_INTEGRATION.md','RES/firmware/docs/SI4463_LIVE_DEBUG_20260924.md','RES/firmware/docs/BENCH_NO_CAN_20260924.md','RES/firmware/dbc/RES_CAN_Validation.dbc',
'全车原理图/RZS_TestBench/电池箱内部原理图.pdf','全车原理图/RZS_TestBench/充电安全回路.pdf','与其他组/整车线路9.26.pdf','2025电检结果/2025赛季电检结果.pdf']
records=[]
for name in files:
 p=root/name
 if not p.exists(): continue
 rec={'path':name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
 if p.suffix=='.xlsx':
  w=load_workbook(p,data_only=True); rec['sheets']={}
  for s in w:
   rec['sheets'][s.title]=[{'row':i,'cells':[(str(v) if v is not None else '') for v in row]} for i,row in enumerate(s.values,1) if any(v is not None for v in row)]
 elif p.suffix=='.docx':
  d=Document(p); rec['paragraphs']=[p.text for p in d.paragraphs if p.text.strip()];rec['tables']=[[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
 elif p.suffix=='.pdf':
  rec['pages']=[{'page':i,'text':p.extract_text()} for i,p in enumerate(PdfReader(p).pages,1)]
 else: rec['text']=p.read_text(encoding='utf-8-sig')
 records.append(rec)
(out/'sources.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
for rec in records:
 lines=['SOURCE '+rec['path']]
 if 'sheets' in rec:
  for name,rows in rec['sheets'].items():
   lines.append('SHEET '+name)
   lines += ['R'+str(r['row'])+' | '+' | '.join(r['cells']) for r in rows]
 elif 'tables' in rec:
  lines+=rec['paragraphs']
  for t in rec['tables']: lines += [' | '.join(r) for r in t]
 elif 'pages' in rec: lines += ['PAGE '+str(p['page'])+' '+p['text'] for p in rec['pages']]
 else: lines.append(rec['text'])
 (out/(Path(rec['path']).stem+'.txt')).write_text('\n'.join(lines),encoding='utf-8')
 print(rec['path'], sum(len(v) for v in rec.get('sheets',{}).values()) if 'sheets' in rec else sum(len(x) for x in lines))
