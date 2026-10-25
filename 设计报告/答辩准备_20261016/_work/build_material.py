from pathlib import Path
import re,json,hashlib
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r'D:\git\RZS_Repository'); OUT=Path(__file__).parent.parent
sources=json.loads((OUT/'_work/sources.json').read_text(encoding='utf8'))
sources=[r for r in sources if not r['path'].startswith('VCU/') and r['path']!='RES/firmware/docs/VCU_INTEGRATION.md']
SID={r['path']:f'S{i:02d}' for i,r in enumerate(sources,1)}
def ref(*paths): return '依据：'+'、'.join('['+SID[p]+']' for p in paths if p in SID)+'。'
base=sources[0]['sheets']
external=[r for r in base['外部接口CN7-CN12'] if r['row']>=5]
headers=[r for r in base['板卡排针H1-H20'] if r['row']>=5]
def nets(path):
 t=(ROOT/path).read_text(encoding='utf-8-sig').split('$NETS',1)[1].split('$SCHEDULE',1)[0]
 d={}
 for net,nodes in re.findall(r"^('[^']+'|[^\s';]+)\s*;\s*(.*?)(?=^(?:'[^']+'|[^\s';]+)\s*;|\Z)",t,re.S|re.M):
  name=net.strip().strip("'")
  for node in re.findall(r'\b[A-Z][A-Z0-9]*\.\d+\b',nodes): d[node]=name
 return d
allnets={p:nets('网表/'+p+'.tel') for p in ['base','bspd','ebs','lock','relay','tsal','tsal_light']}
checks=[]
for r in external+headers:
 c=r['cells']; pin=c[0]+'.'+c[1]; signal=c[2]; net=allnets['base'].get(pin,'NC')
 checks.append({'pin':pin,'table_signal':signal,'netlist_net':net,'listed':pin in allnets['base'],'source_row':r['row']})
assert len(external)==72 and len(headers)==47
assert allnets['base']['CN7.11']=='ROTATE_OUT' and allnets['base']['CN7.13']=='ROTATE_IN' and allnets['base']['CN7.14']=='ROTATE_VCU' and allnets['base']['CN8.6']=='GND'
assert allnets['base']['CN11.2']==allnets['base']['H9.2']=='RES_ERROR'
assert allnets['relay']['U14.8']=='MAINP_1' and allnets['relay']['U15.3']=='MAINP_ACTUAL'
assert allnets['tsal_light']['CN4.1']=='RED' and allnets['tsal_light']['CN4.4']=='GREEN'
(OUT/'_work/connection_audit.json').write_text(json.dumps({'scope':'逐针存在性与关键连接检查，不代表电路全功能验收','external':72,'headers':47,'records':checks},ensure_ascii=False,indent=2),encoding='utf8')

# 完整索引保留原表逐行信息，同时把冲突与现用定义放在最前。
md=['# RZS 全车接口与信号索引','', '整理日期：2026-10-16。与《RZS 整车电气系统答辩准备材料》配套。', '',
'本索引完整保留两份硬件接口工作簿的非空行，并汇入板卡接口表、当前七份网表的连接器网络和 RES CAN 信号。原表转录是追溯资料；现用口径、冲突项和验收边界见下节。未把历史针脚或有冲突的电压自动认定为最终接线。', '',
'## 现用口径与版本冲突', '',
'- EBS：EBS_PWR/EBS_WORK 主支路 24 V；SDC_IN 经外部 DCDC 为 12 V，RLY2 为 12 V 线圈。旧表的 SDC_IN 24 V 已校正，见 EBS/EBS说明.md。',
'- RES_ERROR：VCU 独立输出，健康 12 V、故障 0 V；先建立健康状态再允许 TSMS 上电。PB13 FAULT_OUT 是另一根 3.3 V 高有效本地故障输出。',
'- 删除的是 EBS/RES 两路都为 1 后才开始监测的等待设计；实际故障锁存仍保留。',
'- RES 现用 G0B1/Si4463 SPI；F103/SX127x、E220 UART 为历史资料，不能混接。',
'- EBS_TRIG 在基板 CN11.3 是外部 12 V 高有效，经 R7/D21/R41 约 4.7 V 钳位后送基板 H4.2；EBS 板 H1.2 是 5 V 逻辑域入口。',
'- 基板 v7 写 0～500 V 映射 1～5 V，而 TSAL 阈值为 0.300/0.743 V；按线性映射 60 V 为 1.48 V，尚未建立一致的传感器到比较器传递函数。',
'- 继电器检测旧文档 U11/CN1 与现网表 U14/U15 为同类接口的位号差异，装配以板版号和现图核对。',
'- 本索引不纳入VCU程序、引脚分配表、旧整车报文表及集成分析；网络名中的VCU字样仅保留现有硬件信号标识。',
'- 开路不是固定 0 V；ACTUAL/# 不是独立物理冗余；RelayBits 是驱动命令，不是触点测量。','']
def esc(s): return str(s).replace('|','／').replace('\n','<br>')
for i,r in enumerate(sources,1):
 if 'sheets' in r:
  md+=['## '+f'S{i:02d} '+r['path'],'']
  for s,rows in r['sheets'].items():
   if not rows or s=='VCU_IO速查': continue
   md+=['### '+s,'','| 源行 | 原始单元格内容 按列顺序保留 |','|---|---|']
   for row in rows:
    cells=['C'+str(j+1)+' '+esc(v) for j,v in enumerate(row['cells']) if v]
    md.append('| '+str(row['row'])+' | '+'<br>'.join(cells)+' |')
   md.append('')
for r in sources:
 if 'tables' in r and r['tables']:
  md+=['## '+SID[r['path']]+' '+r['path'],'']
  for j,t in enumerate(r['tables'],1):
   md+=['### 原文表 '+str(j),'','| 行 | 单元格 按列顺序保留 |','|---|---|']
   md += ['| '+str(k)+' | '+'<br>'.join(esc(c) for c in row if c)+' |' for k,row in enumerate(t,1)]
   md.append('')
for module,d in allnets.items():
 md+=['## 当前网表 '+module+'.tel 连接器逐针网络','','| 针脚 | 当前网络 |','|---|---|']
 items=[(k,v) for k,v in d.items() if re.match(r'^(CN\d+|H\d+)\.',k) or (module=='relay' and re.match(r'^U1[45]\.',k))]
 for k,v in sorted(items,key=lambda t:(re.sub(r'\d','',t[0]),*[int(x) for x in re.findall(r'\d+',t[0])])): md.append('| '+k+' | '+esc(v)+' |')
 md.append('')
md+=['## RES 现用接口与 CAN 字段','','现用系统说明逐针表与状态机如下。文本描述已实现的功能及其验证边界，不能据此认定实车验收完成。','',next(r['text'] for r in sources if r['path']=='RES/firmware/docs/RES_SYSTEM_DESCRIPTION_CN.md').split('## 8 CAN 报文',1)[0],'', '## RES CAN 输出字段','','本节仅保留RES自身的字段定义，不包含外部控制器集成。','','```text','\n'.join(line.split(' VCU')[0] for line in next(r['text'] for r in sources if r['path'].endswith('.dbc')).splitlines() if line.startswith('BO_ ') or line.startswith(' SG_ ')),'```','', '## 资料清单和内容校验','','每条 SHA256 仅标识本次读取的文件内容，不代表文件已经通过设计或试验验收。','']
for r in sources:md+=['- ['+SID[r['path']]+'] '+r['path']+'  SHA256 '+r['sha256']]
(OUT/'全车接口与信号索引.md').write_text('\n'.join(md),encoding='utf8')

doc=Document(); sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.6);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
sec.header_distance=Cm(.7);sec.footer_distance=Cm(.7)
sec.top_margin=Cm(1.55);sec.bottom_margin=Cm(1.45)
styles=doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
styles['Normal'].font.size=Pt(10)
styles['Normal'].paragraph_format.line_spacing=1.16
for st in styles:
 for border in st._element.xpath('.//w:pBdr'):
  border.getparent().remove(border)
styles['Normal'].paragraph_format.space_after=Pt(5)
for n,size in [('Title',23),('Subtitle',11),('Heading 1',17),('Heading 2',12),('Heading 3',10.5)]:styles[n].font.size=Pt(size)
styles['Heading 1'].paragraph_format.space_after=Pt(9)
styles['Heading 2'].paragraph_format.space_before=Pt(8)
styles['Heading 2'].paragraph_format.space_after=Pt(5)
header=sec.header.paragraphs[0];header.text='合肥工业大学睿智车队  整车电气系统答辩准备';header.runs[0].font.size=Pt(8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('2026年10月16日  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for r in footer.runs:r.font.size=Pt(8)
def p(text,style=None):return doc.add_paragraph(text,style)
def h(text):doc.add_heading(text,2)
def page(title):
 if len(doc.paragraphs)>1:doc.add_page_break()
 doc.add_heading(title,1)
def table(head,rows,widths=None,size=8.4,compact=False):
 t=doc.add_table(rows=1,cols=len(head));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths):c.width=Cm(w)
 for c,text in zip(t.rows[0].cells,head):c.text=str(text)
 for row in rows:
  cells=t.add_row().cells
  for c,text in zip(cells,row):c.text=str(text)
 for ri,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if ri==0:
   repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
  for j,c in enumerate(row.cells):
   if widths:c.width=Cm(widths[j])
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcPr=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'243B53' if ri==0 else ('F2F5F8' if ri%2==0 else 'FFFFFF'));tcPr.append(shade)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   tcPr.append(borders)
   margins=OxmlElement('w:tcMar')
   for edge,val in [('top',25 if compact else 55),('bottom',25 if compact else 55),('left',75),('right',75)]:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');margins.append(el)
   tcPr.append(margins)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(0 if compact else 1);para.paragraph_format.space_before=Pt(0 if compact else 1);para.paragraph_format.line_spacing=1.05
    for run in para.runs:
     run.font.size=Pt(size);run.bold=(ri==0);run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
 p('')
 return t
def src(*paths):
 para=p(ref(*paths));para.paragraph_format.space_after=Pt(3)
 for r in para.runs:r.font.size=Pt(8)

doc.add_paragraph('RZS 整车电气系统答辩准备材料','Title')
p('接口整合  功能实现  整车安全性  可维护性','Subtitle')
p('资料基线  2026年10月16日    适用范围  低压盒及与整车控制相关的电气接口')
h('答辩要说明的核心结论')
p('本设计通过基板汇接、独立硬件安全判断、故障锁存、远程急停与状态反馈，把分散板卡组织成可以定位故障、可以分模块测试的整车电气系统。其安全价值在于故障能够撤销运行许可并保持安全状态；维护价值在于同一根信号可以追溯到连接器、调理电路、功能板和检查方法。')
p('已有依据覆盖接口定义、当前网表、RES 源码及部分历史台架记录。整车安全性提升目前以设计机制和限定条件下的测试结果说明；未形成新版本整车响应时间、制动性能或维护工时改善的完整实测证据。')
h('90秒开场稿')
p('我们的工作是把整车安全功能从各块独立电路，整合为有统一接口、有故障保持、有诊断路径的系统。基板集中连接 BSPD、TSAL、错误锁存与 EBS，把 24 V、12 V、5 V和3.3 V信号按功能分开。BSPD 独立检测制动与驱动功率冲突；IMD、BMS、EBS 和 RES 的异常通过硬件锁存保持安全回路断开；RES 对急停、失联、重放和通信异常撤销许可；EBS 利用失电触发制动的机械储能系统执行停车。')
p('可维护性方面，每个接口都明确针脚、方向、电平、有效条件和掉电状态，能够从外部连接器定位到板卡内部。已有数字台架证明 BSPD 的逻辑与部分时序、旧版锁存功能以及隔离台架下的 RES 断链释放；模拟阈值、最终触点拓扑和整车制动响应仍按验收清单关闭。我们用明确的证据范围说明设计能力。')
h('材料使用方法')
p('主文用于讲述与问答；配套《全车接口与信号索引.md》用于逐针查阅。正文 [Sxx] 对应末尾资料清单，索引保留原工作簿行号及内容哈希。原表中的“正常”指接口整理状态，不等于实车试验通过。本报告不涉及VCU程序、引脚分配、报文迁移及软件集成；保留网络名只用于硬件接线追溯。')

page('1 整车功能分工与安全动作路径')
table(['模块','输入与判断','输出与整车作用'],[
('基板','外部线束、电源、VCU命令及板卡信号','汇接四块功能板；继电器/MOS驱动；状态电平转换'),
('BSPD','四路油压和一路电流的模拟比较','制动与驱动冲突持续有效时断开安全回路'),
('错误锁存','IMD、BMS、EBS、RES异常及复位','保持SCIN–SCOUT断开；健康恢复后仍需正确复位'),
('继电器检测与TSAL','主正、主负、预充辅助触点和隔离电压采样','给出高压相关红绿灯状态，抑制不可信正常显示'),
('EBS','安全回路、24V主电源、WDI与触发命令','释放电磁阀电源并由机械/气路执行紧急制动'),
('RES','操作者GO/STOP、无线有效性、本地通信状态','安全继电器命令、START事件、CAN状态与诊断'),
('VCU','整车传感器、AS前提条件、RES状态和反馈','整车状态管理、RES_ERROR硬线和执行器控制')],[2.8,6.9,7.7],9)
h('三条需要分开说明的链路')
p('安全动作：异常输入 → 独立硬件判断或 RES 安全状态机 → 撤销安全回路许可或 EBS 供电 → 牵引断能与制动。最终串联顺序、RES 两路触点如何接入和 DCDC 位置仍需以冻结后的整车原理图核定。')
p('故障保持：异常消失不等于自动恢复运行。锁存板保存故障，复位只恢复具备条件的许可；RES 软件恢复需要新会话及 READY 确认，GO 仍需再次操作。BSPD 的延时自动恢复属于另一套机制。')
p('诊断显示：灯、蜂鸣器、CAN状态和调试变量帮助定位，不能代替触点导通、SDC 断开或制动压力的物理测量。')
src('基板/基板全部接口特性与定义_2026-07-31_v7.xlsx','错误锁存/说明.txt','RES/firmware/docs/RES_SYSTEM_DESCRIPTION_CN.md')

page('2 电压域和接口使用边界')
table(['接口类别','现用定义','实现及边界'],[
('24 V主电源','BSPD +24VIN；EBS_PWR/EBS_WORK；ASMS输入','电源能量支路与逻辑状态分开；保险和负载电流需定额'),
('12 V电源及状态','基板供电；RES_ERROR健康；EBS SDC_IN','EBS安全回路经外部DCDC降压；RES_ERROR故障为0V'),
('5 V逻辑与模拟','EBS_ERROR；VCU_WDI；OIL/CURRENT；ACTUAL/#','WDI为0～5V推挽10Hz；不能一律按12V VCU接口接线'),
('3.3 V逻辑','G0B1 GPIO、SPI、I2C、START_OUT、FAULT_OUT','只接兼容逻辑或模块输入；不直接驱动继电器线圈'),
('低边开关输出','TSAL红绿阴极；错误灯与部分蜂鸣器','导通拉地、关闭高阻；负载正端由匹配电源供给'),
('高边输出','RUN/DI8、制动灯、ASSI；三路PMOS输出','来自12V母线；继电器释放时多为开路，PMOS通道有下拉'),
('无源触点','SCIN/SCOUT、SC_IN/SC_OUT、辅助触点','不产生固定逻辑电压；检测辅助触点禁止外加有源电压')],[3.2,6,8.2],9)
h('三个不能混用的信号')
p('RES_ERROR 是 VCU 输出到锁存板的 12/0 V 健康信号；车载 PB13 FAULT_OUT 是 3.3 V 高有效故障；EBS_ERROR 是锁存板侧 5/0 V 逻辑。它们的来源、极性和用途均不同。EBS AS_LOCK 在24V工作电压下经22kΩ/4.7kΩ分压约为4.23V，不能直接当作24V回路或已冻结的完整5V输出额定。')
h('电平转换与隔离的区别')
p('基板 U15/U19/U20 的光耦输入与输出共用 GND，仅实现电平转换和缓冲。三路 PMOS 的负载能量来自基板12V，24V状态只控制开关；它们不能被描述为24V转12V电源或电气隔离。')
h('TSAL量程冲突需要关闭')
p('基板表给出0～500V映射1～5V，按线性关系60V对应1.48V；TSAL板参考却为0.300V和约0.743V。除非另有已确认的比例变换，否则无法据此证明60V判定正确。答辩应展示完整传感器传递函数、阈值和实测翻转点，当前列为高优先级待确认。')
src('基板/基板全部接口特性与定义_2026-07-31_v7.xlsx','TSAL/定义.md','EBS/EBS说明.md')

# 全部72个外部针脚按连接器分组，逐行给出最直接的工作定义。
for group,title in [(['CN7'],'3 基板外部接口 CN7'),(['CN8'],'4 基板外部接口 CN8'),(['CN9','CN10'],'5 基板外部接口 CN9和CN10'),(['CN11','CN12'],'6 基板外部接口 CN11和CN12')]:
 page(title)
 p('方向以基板及其所连接功能板整体为观察点。针脚、电气结构与条件由基板v7表提取；完整内部处理、故障状态、状态标注和源行号见配套索引。')
 rows=[]
 for r in external:
  c=r['cells']
  if c[0] not in group:continue
  signal=c[2]; electrical=c[5];beh=c[6]
  if signal=='GND':electrical='0V公共低压地';beh='作为低压回流；核对实体接地'
  if signal=='NC':electrical='未连接';beh='线束保留，不作为地或信号'
  if c[0]=='CN11' and c[1]=='2':electrical='健康12V／故障0V';beh='VCU输出，经H9.2送锁存；先于TSMS建立健康态'
  if c[0]=='CN11' and c[1]=='10':electrical='经DCDC后12V';beh='直通H3.1送EBS；覆盖旧表24V写法'
  if c[0]=='CN9' and signal.startswith('OIL'):beh='约0.75V制动阈值；开路由BSPD上拉处理，须做故障注入'
  if signal=='CURRENT':beh='约1.11V门限；450V假设下对应约5kW'
  electrical=electrical if len(electrical)<=22 else electrical[:20]+'…'
  beh=beh if len(beh)<=34 else beh[:32]+'…'
  rows.append((c[0]+'.'+c[1],signal,electrical,beh))
 table(['针脚','信号','电平或类型','作用与有效条件'],rows,[1.5,4.1,4.0,7.8],8.5,compact=True)
 if group==['CN7']:p('CN7.15/17/18直接提供12V线圈驱动电流，约33mA；CN7.16/19/20是高阻逻辑命令。不能用同一类MCU GPIO驱动假设覆盖两者。')
 if group==['CN8']:p('AS_PWR由外部提供，基板不生成该电源。RESET采用开漏/开集拉低；约4.7V共用偏置与锁存板5V上拉共同存在，需验证上电、断电和复位脉冲。')
 if group==['CN11','CN12']:p('EBS_TRIG：CN11.3为外部12V，经过基板限流钳位后才进入EBS逻辑域。EBS_CHECK与VCU_FB的极性不能由网络名推断，需正常、触发、WDI停止、掉电及断线真值表。')
 src('基板/基板全部接口特性与定义_2026-07-31_v7.xlsx','EBS/EBS说明.md','错误锁存/说明.txt')

for idx,(start,end) in enumerate([(1,10),(11,20)]):
 page(f'{7+idx} 基板板卡排针 H{start}至H{end}')
 p('本页方向相对所连接功能板；这里的 H 编号属于基板。功能板本身的 H 编号是另一套位号，必须连同板名阅读。')
 rows=[]
 for r in headers:
  c=r['cells'];n=int(c[0][1:])
  if not start<=n<=end:continue
  electrical=c[5]
  if c[0]=='H3' and c[1]=='1':electrical='12V安全回路输入，按10月方案'
  if c[0]=='H9' and c[1]=='2':electrical='健康12V／故障0V'
  rows.append((c[0]+'.'+c[1],c[2],electrical,c[6]))
 table(['基板针脚','信号','电气定义','来源或去向'],rows,[1.7,4,5,6.7],8.4)
 p('关键路径：BSPD安全回路输出经基板H6.2 → H19.1进入错误锁存；锁存输出H19.2 → CN11.6 SCOUT。该路径是网表连接依据，不能取代整车其余串联元件的完整图纸。')
 if idx==1:p('H17.3 R_E_LED为单节点预留，未接K12；EBS/RES错误灯由CN7.16命令驱动，不能声称已有独立硬件R_E_LED报警路径。')
 src('基板/基板全部接口特性与定义_2026-07-31_v7.xlsx')

page('9 BSPD 制动与驱动冲突检测')
p('BSPD用于识别持续制动同时仍有较高驱动功率的冲突，采用比较器、门电路、RC计时和常开继电器，安全判断不依赖VCU运行。四路油压采用OR，再与电流门限结果进行AND，避免只观察单个制动通道。')
table(['板卡针脚','信号','定义'],[('H1.1／H1.2','CURRENT／GND','0～50A映射0～5V；电流门限约1.11V'),('H2.1／.2／.3／.4','OIL4／OIL3／OIL2／OIL1','0～20MPa映射0～5V；30bar约0.75V'),('H3.1／H3.2','SC_IN／SC_OUT','常开继电器无源触点；故障或失电断开'),('H4.1／H4.2','+24VIN／GND','24V供电，经LM2596产生内部5V')],[3.5,5,8.9],9)
h('功能如何实现')
p('油压比较参考R27/R28为5.62kΩ/1kΩ，计算约0.755V；电流参考R25/R26为34.8kΩ/10kΩ，计算约1.116V。油压比较结果经U4级联OR，和电流比较结果构成冲突。R8/C11对持续时间进行判定，触发后释放继电器；冲突消失后由R10/C3保持一段恢复时间。')
p('当前用固定电流门限近似5kW，前提是牵引电压450V。按11.16A计算，在400V时对应约4.46kW，500V时约5.58kW；这是本次根据电路参数作的推算，不能表述为实时测量电压与电流后计算功率。')
h('失效与恢复')
p('输入各有10kΩ上拉，开路趋向高电平；单根油压开路不必然单独触发，仍需电流条件，故完整故障安全性要按组合验证。BSPD掉电使触点断开。RC自动恢复只恢复BSPD许可，整车能否再次运行仍取决于锁存、VCU状态和重新启动条件。')
h('可用于答辩的实测')
p('2026年7月数字台架：四路油压OR与电流AND有效；100ms冲突未触发；持续冲突约0.22s断开；清除后约14.1～14.2s恢复；持续冲突15s保持断开；再次有效冲突重启恢复计时；24V掉电触点断开。测量条件为0/5V数字刺激，不代表0.755V/1.116V模拟阈值已通过。')
src('BSPD/BSPD定义.md','ARDUINO测试/测试汇总.md')

page('10 错误锁存与受控复位')
p('锁存板将IMD、BMS分别保持，并将EBS/RES合并为R_E故障通道。三路状态健康才允许继电器闭合；异常恢复后仍须相应复位，阻止短时故障消失后车辆自行重新获得运行许可。')
table(['锁存板针脚','信号','电气与功能'],[('H1.1','IMD_ERROR','健康约12V，故障约9V，比较器判定'),('H1.2','BMS_ERROR','健康约12V，故障约10V，比较器判定'),('H2.1','EBS_ERROR','健康约5V，故障0V'),('H2.2','RES_ERROR','健康约12V，故障0V；20kΩ/12kΩ分压'),('H3.1／.2／.3','IMD_RESET／BMS_RESET／R_E_RESET','板内5V上拉，外部开漏拉低；按通道清除'),('H4.1／H4.2','SCIN／SCOUT','安全回路无源触点'),('H5.1／.2／.3','IMD_LED／BMS_LED／R_E_LED','5V高边状态输出，关闭可能悬空'),('H6.1／H6.2','12V／GND','内部5V稳压，TPS3808上电复位')],[2.7,5.8,8.9],9)
h('硬件实现与安全收益')
p('状态与逻辑满足RELAY_ENABLE = IMD_STATE AND BMS_STATE AND R_E_STATE。比较电路将输入状态电压转换为逻辑，触发器保存故障，驱动级控制常开继电器。当前网表U11的D接5V、U9故障输出送CLK，复位与POR参与清除路径。该硬件保持不能由RES通信恢复代替。')
p('RES分压后：12V约为4.5V，而5V仅约1.875V，因此把RES健康输入误当作5V会造成不可靠高电平。IMD/BMS也不能用统一“故障0V”描述，必须核对真实设备输出窗口。')
h('上电顺序与复位边界')
p('已删除“两路均为1后才开始监测”的等待使能。现用约定是在TSMS上电前建立RES_ERROR=12V。VCU启动、无有效状态或自检失败时仍应故障输出0V并保持TSMS关闭；禁止仅为启动方便强行输出健康。复位必须在故障条件解除后执行，不能在故障持续时解除保护。')
p('7月旧台架记录通过四路故障保持、故障中防复位、恢复后相应复位和复位隔离；9月另有板级排查，且10月上电顺序变更未重新验收，因此这些历史通过不代表现板现配置已经完成整体验证。')
src('错误锁存/说明.txt','ARDUINO测试/测试汇总.md')

page('11 EBS 失电制动与VCU运行监督')
p('EBS电气部分控制电磁阀能量许可，制动能量由机械储能及气液系统提供。设计目标是供电丢失或控制异常直接触发制动；电路只能证明电源切断路径，实际建压、减速和保持需要气路及整车试验。')
table(['EBS板针脚','信号','现用接口与处理'],[('H1.1','VCU_WDI','VCU输出0～5V推挽10Hz；直接进入5V看门狗'),('H1.2','EBS_TRIG','5V逻辑域高有效；外部12V命令先经基板调理'),('H1.3','EBS_CHECK','检查反馈；外部门限、极性、源阻抗待测'),('H2.1','SDC_IN','外部DCDC后12V，驱动RLY2 12V线圈'),('H2.2','EBS_WORK','24V电磁阀工作电源支路，释放后失电'),('H2.3','AS_LOCK','24V经22kΩ/4.7kΩ分压约4.23V'),('H2.4','VCU_FB','继电器反馈，极性与幅值按状态组合核对'),('H3.1／H3.2','EBS_PWR／GND','24V主电源，LM2596产生5V；RLY1为24V')],[2.6,3.6,11.2],9)
h('实现路径')
p('24V主电源 → EBS_PWR → RLY1触点 → EBS_WORK；安全回路 → 外部DCDC 12V → SDC_IN → RLY2线圈。两只继电器控制不同支路，不能将安全回路降压解释为整个EBS改为12V。WDI监督及EBS_TRIG参与RLY1控制，失电使电磁阀撤销保持，机械气路执行制动。')
h('两项边界必须讲清')
p('旧EBS功能说明有“上电解除刹车、复位解除刹车”等概括，不能据此允许电气复位自动解除已经激活的机械保持。电气许可、机械锁存、人工释放和再次AS Ready应分开验证。气压异常触发、备用电容供电等旧文字尚无完整现网表到实物闭环，答辩中作为系统需求或待核对项，不列为已完成电路能力。')
p('WDI存在只能说明有翻转，不能单独证明VCU所有安全任务健康。需覆盖停止、常高、常低、错误频率及安全任务卡死而输出仍翻转的场景。EBS实际触发到减速时间不能从继电器或GPIO动作时间推定。')
src('EBS/EBS说明.md','EBS/EBS功能说明（2026）.docx','基板/基板全部接口特性与定义_2026-07-31_v7.xlsx')

page('12 继电器机械状态检测')
p('检测板位于电池箱内，采集主正、主负、预充继电器的无源辅助触点。每只继电器用一对触点导线连接检测板；不接主功率端，不接线圈，也不注入外部有源电压。')
table(['现网表针脚','网络','接线用途'],[('U14.1／.2','GND／12V','供电；旧文档对应U11.1／.2'),('U14.3／.4','PRE_2／PRE_1','预充辅助触点两端'),('U14.5／.6','MAINN_2／MAINN_1','主负辅助触点两端'),('U14.7／.8','MAINP_2／MAINP_1','主正辅助触点两端'),('U15.1','PRE_ACTUAL','→基板CN9.5→H18.2→TSAL H7.2'),('U15.2','MAINN_ACTUAL','→CN9.1→H20.1→TSAL H6.1'),('U15.3','MAINP_ACTUAL','→CN9.9→H16.2→TSAL H5.2'),('U15.4','PRE#','→CN9.3→H18.1→TSAL H7.1'),('U15.5','MAINN#','→CN9.2→H20.2→TSAL H6.2'),('U15.6','MAINP#','→CN9.7→H16.1→TSAL H5.1')],[3,4.3,10.1],9)
h('采集如何实现')
p('每路_1端由4.7kΩ上拉到5V；触点闭合把高电平送到_2，断开时_2由100kΩ下拉到低，并用100nF滤波。缓冲和反相器生成ACTUAL与#互补编码；基板六路串47Ω后送TSAL。')
table(['辅助触点','ACTUAL','#','解释'],[('断开','0','1','当前检测链报告断开'),('闭合','1','0','当前检测链报告闭合'),('任意','0','0','编码不一致，不判可信正常'),('任意','1','1','编码不一致，不判可信正常')],[4,2,2,9.4],9)
h('安全与维护收益')
p('监测的是辅助触点状态，比仅观察线圈命令更接近机械状态；但辅助触点和主触点的对应关系要由器件规格与试验确认。ACTUAL/#来自同一物理节点，可以校核传输和编码一致性，不能宣称两路独立传感器或已能覆盖所有主触点粘连。')
p('旧DOCX使用U11/CN1，当前relay.tel使用U14/U15。针脚功能顺序相符，但实板应以板版号和丝印逐针核对，此处不把历史位号覆盖到实物。')
src('继电器检测/继电器检测说明.docx','TSAL/定义.md','基板/基板全部接口特性与定义_2026-07-31_v7.xlsx')

page('13 TSAL 高压状态显示与灯板')
p('TSAL把电压状态与继电器机械状态结合，分别驱动红灯闪烁和绿灯显示。其安全价值是减少仅凭控制命令判断高压状态的误导；当前尚未完成全部传感器阈值和带负载显示验收。')
table(['TSAL板接口','信号','作用'],[('H1.1／.2','NC／V_CONTROLLER','控制器电压低压采样；红灯参考约0.743V'),('H2.1／.2','NC／V_BATTERY','电池电压低压采样；参考约0.300V'),('H3.1／.2','GND／12V','两路内部5V的共同上游电源'),('H4.1／.2','GREEN_OUT_K／RED_OUT_K','低边输出；有效拉地，关闭高阻'),('H5.1／.2','MAINP#／MAINP_ACTUAL','主正辅助触点互补编码'),('H6.1／.2','MAINN_ACTUAL／MAINN#','主负辅助触点互补编码'),('H7.1／.2','PRE#／PRE_ACTUAL','预充辅助触点互补编码')],[3,5.4,9],9)
h('功能实现')
p('红灯由TL331比较V_CONTROLLER后控制NE555，再由Q1低边驱动闪烁。绿灯由LM393、电平组合、断线锁存与传感器就绪控制构成，Q2驱动绿灯阴极。TPS3840及定时电容参与启动/就绪；两路LM340仍共享12V和地，不能表述为完整电源独立。')
table(['灯板CN4针脚','网络','连接'],[('1','RED','接RED_OUT_K低边红灯驱动'),('2','GND','低压地'),('3','12VIN','灯板12V电源'),('4','GREEN','接GREEN_OUT_K低边绿灯驱动')],[3,4,10.4],9)
p('灯板现网表有两组红灯串与两组绿灯串及串联电阻，红绿阴极分别引出。输出允许电流、灯板温升、亮度与闪烁频率仍需实物测量；不把器件额定直接当作整个输出接口额定。')
h('验收和答辩边界')
p('应按实际灯负载扫描两路电压输入、完整触点组合及00/11失配，注入断线/短路并观察上电、掉电和回灌。历史台架只完成部分数字逻辑观察；无外部负载时用Nano内部上拉得出的“都亮”读数不能作为灯状态证据。电压传感器量程冲突见第2节，关闭后才可给出60V对应点。')
src('TSAL/定义.md','ARDUINO测试/测试汇总.md')

page('14 RES 双端硬件与外部接口')
p('现用方案为两端STM32G0B1CBT6配S017 Si4463 SPI无线模块。遥控端处理GO/急停、电量和四灯；车载端处理安全继电器、START、FAULT及CAN。旧F103/LoRa和E220 UART表只留作历史追溯。')
table(['端点或功能','MCU针脚','外设与电气定义'],[('双端S017片选／时钟','PA4／PA5','nSEL脚9低有效／SCLK脚6'),('双端S017数据','PA6／PA7','SDO脚7→MISO／MOSI→SDI脚8'),('双端S017中断／关断','PB2／PB3','nIRQ脚5低有效／SDN脚10高关断'),('双端S017供电','3.3V／GND','VCC脚1和2；GND脚11和12均连接'),('遥控GO','PB0','低有效，NO按钮接地，25ms消抖'),('遥控STOP','PB1','健康NC检测链拉低；高表示急停、断线或检测链掉电'),('遥控电量I2C','PB6／PB7','INA226 SCL／SDA；3.3V上拉；7位地址0x40'),('遥控电量告警','PA8','ALERT保留，当前主要I2C轮询'),('遥控四灯','PB13／PB12／PB11／PB10','IN1红／IN2绿／IN3黄／IN4蓝；低有效'),('车载安全继电器','PB10／PB11','R1／R2低有效；复位/故障释放须实测'),('车载启动指示','PB12','START_OUT 3.3V高有效约100ms'),('车载故障指示','PB13','FAULT_OUT 3.3V高有效；不是RES_ERROR'),('车载CAN逻辑','PB9／PB8／PA0','TXD／RXD／STB；须接收发器'),('双端调试','PA13／PA14','SWDIO／SWCLK，各自独立探针与目标身份')],[3.5,4.0,9.9],8.5)
h('电源与布线')
p('Si4463只能由稳定3.3V供电，同端模块与MCU共地。SPI信号尽量短，远离继电器线圈、DC/DC输入和天线馈线。继电器模块逻辑须兼容3.3V；MCU引脚不能承载线圈。硬件上拉、供电顺序和复位状态需要覆盖MCU无电、下载和看门狗复位场景。')
p('INA226按10mΩ分流器标定，预警11.1V、严重低电量10.5V；采样无效应阻止正常许可，不能伪造正常读数。CAN使用500kbit/s Classic CAN，总线两端各120Ω。CANB与CAN1的命名仍需结合VCU最终线束核对，不能仅凭名称认定物理端口。')
src('RES/firmware/docs/RES_SYSTEM_DESCRIPTION_CN.md','RES/firmware/docs/SI4463_LIVE_DEBUG_20260924.md')

page('15 RES 安全状态机与通信保护')
table(['状态或事件','R1命令','R2命令','动作与恢复条件'],[('复位和初始化','断开','断开','START无效；未知状态不允许运行'),('WAIT_LINK','闭合','断开','仅车载待机，不能解释为已获整车许可'),('有效新会话及3次READY','闭合','闭合','只恢复RES许可，不产生GO'),('新GO事件','闭合','闭合','一次START脉冲及CAN事件；相同命令重发不重复启动'),('STOP／失联／本地故障','断开','断开','取消START并锁存故障'),('普通通信恢复','保持停止','保持停止','同会话不自动解除停止'),('遥控重启且条件满足','按状态恢复','按状态恢复','新会话再确认；本地故障仍禁止恢复')],[5,2,2,8.4],9)
h('通信如何避免错误启动')
p('协议包含配对标识、CRC、鉴权、启动会话、单调序号和重放检查。CRC处理传输错误，鉴权处理报文来源，会话与计数器防止旧命令重放。READY、GO、STOP分别表达许可、一次发车事件和撤销许可；STOP优先级最高。演示密钥不能作为正式每车独立密钥。')
p('现用默认车载代码无线超时为500ms，状态周期50ms，START脉冲100ms，READY确认次数3。遥控默认心跳200ms、GO消抖25ms；工程另有专用台架宏可改变时间，核对实际构建参数才能把源码参数关联到实物映像。')
h('软件恢复与整车硬件复位分开')
p('遥控端重启建立新会话，可在本地故障解除后恢复RES软件许可；它不会清除整车非可编程锁存，也不会替代EBS人工机械释放。恢复后必须再次给出新的GO；单纯网络恢复、按钮松开或重复报文均不得自动启动车辆。')
h('继电器拓扑的现存差异')
p('资料同时存在“R1切EBS供电、R2切SDC”和“两路常开串联SDC”两种描述。现文保留差异，未擅自选定最终接法。不能据两路GPIO就声称已有两条独立物理断开通路或完整单点故障容错；需按最终图纸与触点实测确定。')
src('RES/firmware/docs/RES_SYSTEM_DESCRIPTION_CN.md','RES/firmware/docs/VCU_INTEGRATION.md')
p('补充实现依据：RES/firmware/vehicle_g0b1/Core/Src/res_vehicle_app.c；remote_g0b1/Core/Src/res_remote_app.c。已读取默认宏，未在本次修改或烧录。')

page('16 CAN契约与VCU安全监督')
table(['报文','周期或触发','字段与处理'],[('0x510 RES_STATUS','50ms，8字节标准帧','B0版本；B1状态；B2～3故障标志；B4信道；B5继电器命令；B6～7 GO事件计数，小端'),('0x511 RES_GO_EVENT','每次新GO事件','B0～3遥控会话；B4～7命令计数，均32位小端'),('0x512 RES_DIAGNOSTIC','随GO事件发送','B0～3命令计数；B4～5事件计数；B6信道；B7版本')],[4.1,4,9.3],9)
h('状态与故障字段')
p('0x510状态：0 BOOT，1 WAIT_LINK，2 READY，3 GO_EVENT，4 STOPPED_FAULT。B2的bit0～7依次是链路有效、远端停止、无线超时、鉴权失败、继电器不一致、会话恢复、CAN本地故障、软件锁存。B3保留；B5的bit0/1为两路继电器命令，其余保留。')
p('RelayBits表达输出命令，不等于触点闭合反馈。DBC中的relay_mismatch标志只能说明协议保留了该故障项，不能证明已安装独立触点采集并可发现所有粘连。')
h('VCU端必须实现的监督')
p('在固定周期安全任务检查最后有效0x510距今是否超过150ms，并检查版本、状态、故障标志、READY下两路命令一致性与独立AS Ready前提。总线静默时接收回调不再执行，因此不能只在回调中检查超时。CAN bus-off、解码异常或安全任务超时应驱动RES_ERROR=0V。')
p('接受0x511还需验证新命令计数、最近READY且链路有效、两路许可命令有效和整车启动条件。START_OUT仅为约100ms电气指示，不得绕过这些检查。VCU监督文件提供了示例逻辑，仍需核对VCU实际生成代码及接线，不能把接口契约直接写成VCU已集成完成。')
h('历史报文迁移')
p('旧整车报文表为res_msg=0x1E4、50ms。新RES契约使用0x510/0x511/0x512，二者不是同一帧的别名。必须同步VCU解码器、DBC、总线端口、节点配置和测试报文。旧表ESP、IPC、AMI、EPS及电机ID保留在信号索引，未在本次重定义其协议。')
src('RES/firmware/dbc/RES_CAN_Validation.dbc','RES/firmware/docs/VCU_INTEGRATION.md','VCU/方程式/方程式/报文id(1)(1).xlsx')

page('17 VCU接口映射与系统功能')
p('VCU工作簿同时包含软件功能表、硬件121PIN接口表和另一套外部引脚示意。下表列出明确标注的关键用途，作为联合核对入口；不作为已冻结的整车线束制造表。额定电流为源表标注，未作本轮带载验证。')
table(['功能','源表接口','作用与基板适配'],[('踏板和制动采样','AIN0 PIN51；AIN1 PIN70；AIN2 PIN50','分别加速踏板1/2及制动踏板；最大5V输入'),('气压1/2','AIN4 PIN49；AIN5 PIN68','表标最大10V，实际传感器量程和软件标定须一致'),('轮速前/后','AIN6 PIN48；AIN7 PIN67','功能表称霍尔，硬件表列模拟量；采集方式需核对'),('EBS状态与急停起止','HIN1 PIN59；HIN2 PIN78；HIN3 PIN20','源表为高边输入；基板反馈幅值和故障极性另测'),('TSMS／ASMS','HIN7 PIN19；HIN8 PIN38','VCU状态机输入；ASMS基板输出为12V'),('制动灯／ASSI','OUT1 PIN53；OUT2 PIN34；OUT3 PIN15','低边输出适配基板低有效线圈控制'),('EBS失效灯','OUT4 PIN73','源表低边，与基板CN7.16的12V高有效命令冲突，需中间驱动/重分配'),('RES错误','OUT5 PIN54 resl_err','源表低边；不等于已具备独立健康12V硬线，须核对上拉/高边驱动'),('EBS触发','HO1 PIN31','源表高边输出，基板CN11.3外部12V输入'),('BMS／IMD故障灯','HO2 PIN12；HO3 PIN11','对应基板CN7.19／20高有效命令，电压及故障默认态需核对'),('DI8／RUN／R2D','HO4 PIN30；HO5 PIN27；HO6 PIN8','对应基板CN7.17／18／15，需源出约33mA'),('WDI','PA30 PIN108','源表5V/PWM预留并标看门狗，需输出10Hz推挽方波'),('CAN0／CAN1／CAN2','PIN45/64；46/65；47/66','源表写L/H组合；协议与整车CANB对应关系需核定')],[3,5.4,9],8.2)
h('软件功能分工')
p('Sensor_Signal采集踏板、轮速和气压；ASB_Selfcheck检查EBS相关可靠性；TS_Activation与Motor_Control管理驱动激活和电机命令；AS_Stateflow管理AS状态；IPC/EPS/ESP/AMI负责相应通信；LED_Control和Buzzer负责提示。功能表中的0/5V或uint8逻辑量属于软件表示，不自动等于连接器实际电压与高低边驱动能力。')
p('NXPS32K“引脚分配图”还包含另一套PIN24等数字输入与CAN PIN33/14等映射，与“VCU接插件示意图”不一致；需按实际VCU型号、板版号和线束连续性选择正确版本。')
src('VCU/方程式/方程式/NXPS32K_引脚分配(1).xlsx','VCU/方程式/方程式/VCU部分引脚改(1).xlsx','基板/基板全部接口特性与定义_2026-07-31_v7.xlsx')

page('18 整车安全性提升的论证')
p('答辩建议按“危险来源 → 检测机制 → 安全动作 → 故障保持 → 验证证据”讲述。以下是由现有设计推导的安全收益，不包含尚无试验支持的事故率或失效率改善数字。')
table(['风险场景','设计应对','整车安全收益与验证边界'],[('制动时驱动仍输出','BSPD模拟比较、四路OR、冲突AND和延时','独立于VCU撤销许可；数字时序有旧台架依据，真实阈值与全压范围待测'),('绝缘或电池故障短时消失','IMD/BMS故障硬件保持','防止恢复即重启；真实设备输出窗口与现板需复验'),('远程急停或无线丢失','STOP优先、500ms超时、两路释放命令','链路失效转向停止；实际SDC时间与射频可靠性待测'),('重复或重放GO','鉴权、会话、序号、命令计数','减少无新操作却重复启动；协议测试不替代安全评估'),('VCU或CAN失效','WDI监督；RES状态150ms契约','多种故障入口撤销许可；VCU集成、WDI覆盖率待核定'),('失电和MCU复位','常开安全触点；复位默认释放；BOR/IWDG机制','能量丢失倾向停止；供电交错和引脚回灌需验证'),('状态显示不可信','电压检测加辅助触点及互补一致性','降低仅凭命令显示安全的风险；非独立传感冗余'),('误接电平或错误复位','明确电压域、调理和开漏复位','减少12V直入5V逻辑及非预期清锁；仍需逐针检查')],[3.1,5.5,8.8],9)
h('不要夸大独立性和冗余')
p('硬件BSPD与VCU控制分工可以减少对VCU软件的共同依赖，但共用供电、线束和地仍可能产生共同原因失效。两路RES驱动命令、两路互补检测和红绿两路稳压分别有不同边界，不能统一称为“双冗余”。')
h('整车安全状态必须由物理量确认')
p('需要共同确认牵引许可撤销、安全回路断开、EBS制动压力建立、车辆停止并保持。GPIO翻转、继电器声音或CAN故障位只能作为中间证据；完整合格判定应关联实际触点、执行器与车辆表现。')
src('BSPD/BSPD定义.md','错误锁存/说明.txt','TSAL/定义.md','RES/firmware/docs/RES_SYSTEM_DESCRIPTION_CN.md')

page('19 可维护性提升与排障方法')
p('基板将多块安全电路集中固定和汇接，使外部线束与板内跳接有统一边界。接口表把电源、命令、反馈、触点和模拟量分开，维修人员可以按信号路径定位，而不只依赖个人对历史接线的记忆。结构上有利于减少松散板间连线、插错和碰撞；尚无改版前后工时和故障率实测，不能量化宣称。')
table(['改进机制','维修人员如何使用','可证明的收益'],[('逐针接口与统一信号名','从CN针脚查到H排针、功能板及网表网络','减少跨文档查找歧义；119个基板针脚已纳入'),('功能板分工','先隔离电源、信号，再以假负载替换单模块','可缩小故障范围；换板后仍须完整安全复验'),('命令与反馈分开','同时测命令GPIO、触点及下游电压','区分软件未命令、驱动异常与触点异常'),('分电压域与明确默认态','先测供电、地和上电顺序，再注入逻辑','降低用错误电平测试导致损坏或误判的风险'),('CAN与只读诊断','记录状态、故障位、会话、计数及超时','可复现通信失效和区分无线/CAN/电量异常'),('文件版本与内容哈希','记录图纸、网表、固件、线束与板版号','将试验结果绑定实际版本，防止旧结论误用')],[3.1,7.2,7.1],9)
h('建议固定的五步排障')
p('第一步记录现象、固件与板版号，保持高能回路隔离。第二步测24V、12V、5V、3.3V以及对应地，排除掉电、欠压和GPIO回灌。第三步在连接器两端测同一信号，区分线束与板内故障。第四步沿比较器/锁存/驱动/触点逐级检查，并观察CAN与诊断量。第五步消除故障后执行正确复位，再验证保护能再次触发，记录原始数据。')
h('三个典型定位例子')
p('RES无法READY：先看INA226有效性和PB1，再看Si4463 part/ready/error及双向ACK，最后看CAN ACK/bus-off；不能只看发射计数。EBS不释放或误制动：分别测SDC_IN 12V、EBS_PWR 24V、WDI、TRIG、两只继电器与EBS_WORK。TSAL异常：核对传感器低压量程、公共地、ACTUAL/#，再以实际灯负载观察开漏输出。')
src('基板/基板全部接口特性与定义_2026-07-31_v7.xlsx','ARDUINO测试/测试汇总.md','RES/firmware/docs/SI4463_LIVE_DEBUG_20260924.md')

page('20 失效分析与验证矩阵')
table(['故障注入','应观察的安全动作','测量点及通过依据'],[('任一路油压与电流持续超门限','BSPD断开并保持规定恢复时间','输入波形、比较器输出、SC_IN–SC_OUT触点；记录最坏响应'),('IMD/BMS异常及恢复','对应锁存；不经正确复位不闭合','真实输出窗口、状态节点、SCIN–SCOUT；故障中复位不能解除'),('EBS_ERROR或RES_ERROR拉低','R_E锁存并撤销安全回路许可','H2输入、R_E_STATE、触点；TSMS前后时序分别测试'),('WDI常高、常低、停止和错频','撤销EBS工作许可','H1.1、WDO#、RLY1、EBS_WORK、实际制动建压'),('安全回路断开／DCDC失电','RLY2释放并执行安全动作','DCDC输入输出、SDC_IN、RLY2触点及制动'),('遥控STOP／无线切断','两路释放，取消START并保持故障','无线最后有效时间、PB10/11、两个COM–NO及实际SDC'),('车载CAN断开／bus-off','车载故障并由VCU监督撤销健康硬线','CAN抓包、VCU安全任务、RES_ERROR、硬件锁存'),('重放GO／旧会话／异常帧','不重复START，不解除STOP','帧记录、会话与命令计数、START波形、事件数'),('继电器触点粘连','其余实际物理路径仍能安全断开或暴露设计缺口','逐只触点强制闭合，测SDC；当前不能先宣称通过'),('传感器及互补线断路/短路','不产生误导的正常显示或错误许可','模拟输入、00/11组合、实际灯负载、锁存及阈值'),('交错上电／掉电／复位','不因回灌或默认电平获得非预期许可','四个电压域、NRST、继电器线圈与触点、START'),('机械气路与整车停车','建立制动并稳定停车保持','压力、触发到减速时间、减速度、停车距离及保持状态')],[4.3,5.7,7.4],8.5)
p('这是依据当前设计编排的验收计划，并非本轮已经执行的试验。先完成无高能台架与电压域确认，再按正式试验方案推进。阈值、时限及机械性能的正式判据应逐项绑定2026规则原文与最终设计配置。')
src('ARDUINO测试/测试汇总.md','RES/firmware/docs/RES_SYSTEM_DESCRIPTION_CN.md','EBS/EBS说明.md')

page('21 已有证据与可陈述的结论')
table(['项目及记录','已有观察','答辩可陈述范围'],[('基板v7与现网表','72个外部针脚和47个排针针脚纳入；CN7转向接口、RES直通及灯板针脚核对','文档/连接层面追溯；未验证载流、温升及实体线束'),('BSPD 2026年7月数字台架','四路OR、AND、约0.22s触点断开、约14s恢复与掉电断开','限定数字刺激下的逻辑和时序；模拟阈值、全故障覆盖待测'),('锁存板7月台架','四类故障保持、防故障中复位和复位隔离记录','仅当时板版与接线；不能覆盖10月启动方案或后续板级异常'),('TSAL 7月台架','部分数字组合和锁存观察；发现测试共地与回灌问题','部分功能验证；无完整真值表、阈值及外灯负载验收'),('RES 9月无线诊断','修复VCC接线后可识别0x4463并有双向ACK；电量I2C接线修复记录','限定通信与供电诊断条件；非距离/抗干扰/零错包验收'),('RES无CAN隔离台架','READY、一次100ms START软件记录；断链两路GPIO释放；用户确认两路COM–NO断开','不测得触点动作时延；无真实SDC、CAN与整车执行器验收'),('恢复正常CAN版记录','下载校验通过；res_bench_no_can=0；未接CAN时安全故障输出关闭','证明当时恢复记录；当前实物映像和带CAN READY/GO需重新确认'),('历史LTspice报告','原版本仿真/模型与波形记录','未按10月方案重跑；不能作为新版全系统通过证据')],[4.2,6.5,6.7],8.5)
h('三个时间数字的正确解释')
p('500ms是RES默认软件无线超时门限；历史约492ms来自非同步状态块读取的时间差，不能当作物理触点响应时间。100ms START来自软件rise/fall记录，尚非示波器测量。BSPD约0.22s是在数字夹具条件下估算冲突成立至触点断开的结果，不能等同EBS触发到车辆减速。')
p('本次只进行了资料读取、关键连接检查和材料制作，未进行上电测试、软件回归、烧录或整车试验。')
src('ARDUINO测试/测试汇总.md','RES/firmware/docs/BENCH_NO_CAN_20260924.md','RES/firmware/docs/SI4463_LIVE_DEBUG_20260924.md','所有板子完整仿真/simulations/simulation_report.md')

page('22 答辩前优先关闭的版本问题')
table(['优先项','发现与暂定处理','关闭证据'],[('EBS SDC电压','旧XLSX及DOCX的24V已校正；采用10月说明12V方案','外部DCDC接线、SDC_IN实测、两继电器型号与实体一致'),('TSAL量程与60V映射','1～5V传感器输出与0.300/0.743V参考未闭环','传感器型号、传递函数、所有调理比例及阈值扫描'),('RES触点最终拓扑','R1切EBS/R2切SDC与双路串SDC两种文字并存','冻结整车原理图、触点连续性和失效注入结果'),('EBS_TRIG接口域','12V外部命令与5V板端逻辑容易混淆','基板限流钳位网络、各节点电压及掉电反灌测试'),('TSMS前健康顺序','删除旧等待后必须先有有效健康状态','多种上电顺序与掉电时序波形；无效状态不强置健康'),('反馈与输入额定','EBS_CHECK/VCU_FB极性及AS_LOCK映射需确定','正常/触发/WDI失效/掉电/断线完整真值表'),('负载与复位额定','三PMOS电流、AS_PWR、RESET脉宽及容差未冻结','保险与热测试、上电/断电/错误复位测试'),('接口位号与旧校核','relay位号与BSPD旧Q2描述已校订；CN7三针按新网表改为ROTATE','现网表U14/U15、Q2.2接GND、CN7.11/13/14接K15的装配核对')],[3.6,7.1,6.7],8.5)
p('已确认的旧定义已同步修正到接口工作簿和相关Word/Markdown文档；原文备份与校订记录随工作文件保存。仍有冲突的实物接口不擅自定稿，答辩可说明文档已校订，未完成的电气问题仍需验证。')
src('低压盒/硬件电路接口.xlsx','基板/基板全部接口特性与定义_2026-07-31_v7.xlsx','EBS/EBS说明.md','RES/firmware/docs/RES_SYSTEM_DESCRIPTION_CN.md')

page('23 八分钟答辩的讲述顺序')
table(['时间','建议展示内容','要回答的问题'],[('0:00～0:45','系统范围与安全目标','系统负责什么，怎样定义安全状态'),('0:45～1:30','功能分工与三条路径','检测、动作、保持、诊断各由谁承担'),('1:30～2:30','基板与电压域','接口如何防误接；板卡如何集成与定位'),('2:30～3:30','BSPD与错误锁存','VCU异常时谁继续判断；故障消失为何不自动恢复'),('3:30～4:30','EBS与RES','断电、失联和通信异常怎样变成制动动作'),('4:30～5:15','继电器检测与TSAL','为何用实际状态，互补编码边界是什么'),('5:15～6:15','安全收益与失效矩阵','覆盖哪些失效，还有哪些共同原因'),('6:15～7:00','维护流程与典型定位','怎样更快找到电源、线束、软件或触点故障'),('7:00～8:00','已有证据与待关闭项','哪些测过、哪些只是设计；下一步如何验收')],[2.3,6,9.1],9)
h('可直接用于收尾的讲稿')
p('这套设计的提升主要体现在安全动作和维修依据都变得明确。对制动与驱动冲突、绝缘与电池异常、远程急停和通信失效，我们分别安排检测、撤销许可和故障保持；对维护，接口可以从外部连接器追到功能板、调理器件和测试点。已有台架结果支持部分逻辑与时序，但我们仍把模拟阈值、最终触点拓扑和整车制动表现作为正式验收项。后续以同一版本的图纸、网表、固件和实车记录完成闭环。')
h('建议携带的证据页')
p('一张最终整车安全回路图；一张24/12/5/3.3V电压域图；一页关键接口对照；BSPD与RES的带时间轴波形；故障注入前后触点与制动压力记录；一页软硬件版本清单。当前缺少的证据应直接标为待测，不能用图标或示意图代替。')
h('图纸与文档准备方式')
p('展示图中每条信号保留与接口表一致的名字，涉及触点时标COM/NO/NC，涉及板卡时同时写板名和连接器位号。只保留与本次实板一致的图纸；历史接口表供问答追溯，避免作为主展示接线依据。')

page('24 高频问答与回答要点')
qas=[
('VCU死机为什么还能安全停车','BSPD和硬件锁存独立完成各自判断，EBS监督WDI，RES也检查自身通信和本地故障。但WDI仍翻转的部分任务失效、共用电源和最终执行路径需专项测试，不能笼统承诺覆盖所有VCU失效。'),
('为什么发生故障后不能直接恢复','输入恢复不说明危险已消除。硬件锁存要求正确复位，RES要求新会话确认且新的GO，EBS机械释放有独立流程，避免网络恢复或上电瞬间自行启动。'),
('EBS为什么同时有12V与24V','24V支路给主电源和电磁阀；12V是安全回路经DCDC后驱动RLY2。两种线圈与两条支路分工不同。'),
('两路继电器是否说明单点失效安全','只有最终触点拓扑、驱动失效模式和粘连注入证明另一物理路径有效后才能作该结论。当前两份拓扑描述仍需统一，命令位不能替代证明。'),
('ACTUAL和反相信号算不算双冗余','来自同一辅助触点节点，属于互补编码一致性检测，可发现部分线路/编码异常，不是两个独立物理传感器。'),
('CAN断线时怎么发现','RES本地可发现发送故障；VCU还需周期安全任务按最后有效状态150ms超时，驱动独立RES_ERROR硬线，不依赖接收回调再次到来。'),
('无线500ms门限是否已满足整车要求','它是软件参数。最后有效报文到真实SDC断开的最坏时间还包含任务调度和触点释放，须测量；现有约492ms状态读取差不是触点实测。'),
('为什么说维护更方便','统一接口和基板使信号可追到针脚、调理、功能板及反馈；按电源→输入→逻辑→驱动→触点定位。暂无维修工时对比，因此不说缩短了某个百分比。'),
('BSPD是否实时计算5kW','当前是固定电流比较，按450V换算。电压变化时等效功率门限变化，需要核定电压范围及规则判据；不能称为实时功率计算。'),
('目前最需要补齐什么','先关闭TSAL量程、EBS电压和触点拓扑等接口问题，再做上电顺序、模拟阈值、触点响应及完整整车制动验收。')]
for q,a in qas:h(q);p(a)

page('25 资料范围与版本追溯')
p('本次整合2份接口工作簿、5份功能或组间协作文档、当前板卡说明、RES系统/协议/台架记录、历史仿真状态以及整车图纸背景，共22项提取源；另读取7份当前网表和RES两端状态机源码。完整工作簿非空行、板卡接口原表、连接器网络和DBC字段保存在配套索引。')
p('主依据优先级为：明确的最新接口约定用于功能和电压口径，现网表用于当前连接，源码用于软件参数；三者有冲突则登记待确认，不能仅凭文件日期选定全部事实。当前网表不证明实板焊接、历史台架不覆盖新版、软件契约不证明VCU已集成。')
table(['资料组','如何使用'],[('基板与低压盒接口表','逐针与电气定义；标出EBS 24V旧项、BSPD旧校核及TSAL量程差异'),('VCU三份工作簿','传感器、执行器、CAN ID及引脚追溯；内部冲突不静默合并'),('板卡说明与网表','解释比较、锁存、触点、驱动及电源路径'),('RES当前系统/代码/DBC','采用G0B1/Si4463及0x510～512；旧F103表仅作历史'),('旧ASF与组间说明','用于理解职责、气路与显示需求；2024版本不直接作为2026实现依据'),('整车PDF及电池箱图','作为布局与系统背景；本轮仅提取文本，未逐线电气复核'),('历史试验与仿真','按当时版本及测试条件引用；不替代现配置验收')],[5,12.4],9)
h('法规与性能数值的引用边界')
p('仓库含2026赛事规则及往年技术检查资料。本材料不据旧DOCX的规则摘录宣布2026合规，也未完成逐条规则审查。RES“最后有效报文至SDC断开小于1.5s”等要求在本材料按现工程文件的验收目标列示；正式答辩提交前应将每项指标对应到2026原文条款及实际测试记录。')
p('2025电检PDF仅有很少可提取文字，不能据此归纳通过项目或改进幅度。所有负载能力、制动性能、射频距离及维修工时均未在本次获得新实测。')

for chunk in range(0,len(sources),14):
 page('26 资料清单' if chunk==0 else '27 资料清单续')
 p('源编号用于正文定位；文件均位于仓库根目录下。内容哈希详见配套索引。')
 rows=[]
 for r in sources[chunk:chunk+14]:
  role='接口原表' if 'sheets' in r else ('历史或功能说明' if 'tables' in r else ('背景图纸文本' if 'pages' in r else '设计或验证依据'))
  rows.append((SID[r['path']],r['path'],role))
 table(['编号','仓库相对路径','用途'],rows,[1.2,13.2,3],8.5)
 if chunk>0:
  h('另行核对的实现文件')
  p('网表/base.tel、bspd.tel、ebs.tel、lock.tel、relay.tel、tsal.tel、tsal_light.tel；RES/firmware/vehicle_g0b1/Core/Src/res_vehicle_app.c；RES/firmware/remote_g0b1/Core/Src/res_remote_app.c。')
  h('本次交付')
  p('主文档用于答辩准备，配套索引用于全量逐针追溯。已按新方案修正两份接口工作簿、相关Word说明及RES现用入口；电路网表和固件未修改。后续设计变化应同步原理图、网表、板版号、线束、源表、固件契约和测试记录。')

# 按用户范围移除VCU程序与接口集成章节。
body=doc._element.body
removing=False
for el in list(body):
 texts=''.join(el.itertext())
 if el.tag==qn('w:p'):
  visible=''.join(el.xpath('.//w:t/text()'))
  if visible.startswith('16 CAN契约') or visible.startswith('17 VCU接口'):
   removing=True
  elif visible.startswith('18 整车安全性'):
   removing=False
 if removing:body.remove(el)
# 删除专属整车控制分析行、问答和来源组。
for t in list(doc.tables):
 for row in list(t.rows):
  txt=' '.join(c.text for c in row.cells)
  if any(s in txt for s in ['VCU三份工作簿','VCU或CAN失效','车载CAN断开／bus-off','整车状态管理、RES_ERROR硬线和执行器控制']):row._tr.getparent().remove(row._tr)
remove_next=False
for para in list(doc.paragraphs):
 txt=para.text
 if remove_next:para._p.getparent().remove(para._p);remove_next=False;continue
 if txt in ['CAN断线时怎么发现','VCU死机为什么还能安全停车']:
  para._p.getparent().remove(para._p);remove_next=True;continue
 if 'VCU监督文件' in txt or 'NXPS32K' in txt:para._p.getparent().remove(para._p)
# 编号连续化，长源表路径和信号名保留原文。
num=0
for para in doc.paragraphs:
 if para.style.name=='Heading 1':
  num+=1;para.text=re.sub(r'^\d+ ',str(num)+' ',para.text)
for para in doc.paragraphs:
 if 'CAN' in para.text and '150ms' in para.text:
  para.text='RES通信接口仅说明本模块自身的状态输出和故障处理；本报告不涉及外部控制器的集成或软件监督实现。'
# 调整叙述中的控制器名称为职责边界，只保留带下划线等现有信号名。
for el in doc._element.body.xpath('.//w:t'):
 if not el.text:continue
 el.text=re.sub(r'(?<![A-Za-z0-9_/])VCU(?![A-Za-z0-9_/])','外部控制器',el.text)

doc.core_properties.title='RZS整车电气系统答辩准备材料'
doc.core_properties.subject='接口信号整合与安全性可维护性论证'
doc.core_properties.author='睿智车队'
dest=OUT/'RZS整车电气系统答辩准备材料.docx';doc.save(dest)
(OUT/'_work/source_manifest.json').write_text(json.dumps([{'id':SID[r['path']],'path':r['path'],'sha256':r['sha256']} for r in sources],ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({'docx':str(dest),'sources':len(sources),'external_pins':len(external),'board_header_pins':len(headers),'signal_index':str(OUT/'全车接口与信号索引.md')},ensure_ascii=False))
