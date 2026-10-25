from pathlib import Path
from openpyxl import load_workbook
from copy import copy
R=Path(r'D:\git\RZS_Repository')
p=R/'基板/基板全部接口特性与定义_2026-07-31_v7.xlsx';w=load_workbook(p)
for s in w:
 for row in s:
  for c in row:
   if isinstance(c.value,str):
    for pin in ['CN7.11 / ','CN7.13 / ','CN7.14 / ']:c.value=c.value.replace(pin,'')
s=w['外部接口CN7-CN12']
for row in s:
 if row[0].value=='CN7' and str(row[1].value) in ['11','13','14']:
  pin=str(row[1].value);name={'11':'ROTATE_OUT','13':'ROTATE_IN','14':'ROTATE_VCU'}[pin];node={'11':'K15.4','13':'K15.3','14':'K15.2'}[pin]
  values=['CN7',int(pin),name,'接口方向按最终执行器定义核定','K15转向继电器接口','触点额定/线圈电压待核定','连接关系以2026-10-04网表/base.tel为准','不再作为GND；工作及失电状态待核对',f'CN7.{pin}直连{node}','禁止按旧表接地；核对K15型号/引脚/实际负载',f'CN7.{pin} / {node}','新网表已接入，功能额定待确认','覆盖2026-07-31旧表GND定义']
  for c,v in zip(row,values):c.value=v
s=w['总览'];s['A2']='连接依据：网表/base.tel（2026-10-04），优先于基板/base.tel旧副本。CN7.11/13/14为ROTATE_OUT/IN/VCU并接K15，不是GND；CN8.6仍为GND。EBS安全回路12V、主支路24V，RES_ERROR健康12V、故障0V。'
s['B9']='继电器触点或MOS关闭后的开路不是固定0V，除非有明确下拉或外部负载。'
s=w['待确认与测试']
for row in s:
 if row[0].value=='P08':
  for c,v in zip(row,['P08','高','CN7.11/13/14新网表为转向接口','网表/base.tel：ROTATE_OUT/ROTATE_IN/ROTATE_VCU接K15；CN8.6仍为GND','按旧表接地可能导致故障','按最新网表与实板核对K15及线束；不改电路','文档已校订，实物待核对']):c.value=v
s=w['依据与器件手册']
for row in s:
 if row[1].value=='base.tel':
  row[2].value='以网表/base.tel 2026-10-04为当前统一入口；基板/base.tel为旧副本'
  row[5].value='新网表新增K15转向继电器及CN7.11/13/14 ROTATE接口；旧接地表述已修正'
w.save(p)

p=R/'设计报告/答辩准备_20261016/_work/build_material.py';t=p.read_text(encoding='utf8')
t=t.replace("sources=json.loads((OUT/'_work/sources.json').read_text(encoding='utf8'))", "sources=json.loads((OUT/'_work/sources.json').read_text(encoding='utf8'))\nsources=[r for r in sources if not r['path'].startswith('VCU/') and r['path']!='RES/firmware/docs/VCU_INTEGRATION.md']")
t=t.replace("'['+SID[p]+']' for p in paths", "'['+SID[p]+']' for p in paths if p in SID")
t=t.replace("assert all(allnets['base'][p]=='GND' for p in ['CN7.11','CN7.13','CN7.14','CN8.6'])", "assert allnets['base']['CN7.11']=='ROTATE_OUT' and allnets['base']['CN7.13']=='ROTATE_IN' and allnets['base']['CN7.14']=='ROTATE_VCU' and allnets['base']['CN8.6']=='GND'")
t=t.replace("if not rows: continue", "if not rows or s=='VCU_IO速查': continue")
t=t.replace('五份接口工作簿','两份硬件接口工作簿').replace('5份接口工作簿、6份功能或组间协作文档','2份接口工作簿、5份功能或组间协作文档').replace('共27项提取源','共22项提取源')
t=t.replace('完整保留五份接口工作簿的非空行','保留两份硬件接口工作簿的非空行，排除控制器软件专页')
t=t.replace('VCU NXPS32K 工作簿两张表内部也有针脚用途冲突；保留原行供追溯，不能直接按混合表制作线束。','本索引不纳入VCU程序、引脚分配表、旧整车报文表及集成分析；网络名中的VCU字样仅保留现有硬件信号标识。')
# 对已校正的旧项更新说明，不把新网表误称为旧GND版。
t=t.replace('关键GND、RES直通及灯板针脚核对','CN7转向接口、RES直通及灯板针脚核对')
t=t.replace("('EBS SDC电压','旧XLSX及DOCX仍写24V；采用10月说明12V方案'", "('EBS SDC电压','旧XLSX及DOCX的24V已校正；采用10月说明12V方案'")
t=t.replace("('VCU输出适配','OUT4低边与CN7.16高有效；resl_err低边与独立12V硬线','驱动电路、上电默认态、有效电平、故障动作实测'),",'')
t=t.replace("('VCU内部引脚冲突','两张表对部分输入和CAN针脚定义不同','准确VCU型号/版号、最终引脚分配、线束逐点通断'),",'')
t=t.replace("('RES报文迁移','旧0x1E4与新0x510～512并存','实际VCU解码、周期安全任务及故障注入抓包'),",'')
t=t.replace("('接口位号与旧校核','relay U11/CN1旧名；BSPD旧Q2方向问题仍在旧表','现网表U14/U15；bspd Q2.2接GND的装配核对')", "('接口位号与旧校核','relay位号与BSPD旧Q2描述已校订；CN7三针按新网表改为ROTATE','现网表U14/U15、Q2.2接GND、CN7.11/13/14接K15的装配核对')")
t=t.replace('原表中的“正常”指接口整理状态，不等于实车试验通过。','原表中的“正常”指接口整理状态，不等于实车试验通过。本报告不涉及VCU程序、引脚分配、报文迁移及软件集成；保留网络名只用于硬件接线追溯。')
# 追加后处理：从文档中移除整车控制器章节与专属分析，保留板卡自身物理接口。
needle="doc.core_properties.title='RZS整车电气系统答辩准备材料'"
post='''# 按用户范围移除VCU程序与接口集成章节。
body=doc._element.body
removing=False
for el in list(body):
 texts=''.join(el.itertext())
 if el.tag==qn('w:p'):
  visible=''.join(el.xpath('.//w:t/text()'))
  if visible.startswith('16 CAN契约') or visible.startswith('17 VCU接口'):
   removing=True
  elif visible.startswith('18 整车安全性'):
   removing=False
 if removing:body.remove(el)
# 删除专属整车控制分析行、问答和来源组。
for t in list(doc.tables):
 for row in list(t.rows):
  txt=' '.join(c.text for c in row.cells)
  if any(s in txt for s in ['VCU三份工作簿','VCU或CAN失效','车载CAN断开／bus-off','整车状态管理、RES_ERROR硬线和执行器控制']):row._tr.getparent().remove(row._tr)
remove_next=False
for para in list(doc.paragraphs):
 txt=para.text
 if remove_next:para._p.getparent().remove(para._p);remove_next=False;continue
 if txt in ['CAN断线时怎么发现','VCU死机为什么还能安全停车']:
  para._p.getparent().remove(para._p);remove_next=True;continue
 if 'VCU监督文件' in txt or 'NXPS32K' in txt:para._p.getparent().remove(para._p)
# 编号连续化，长源表路径和信号名保留原文。
num=0
for para in doc.paragraphs:
 if para.style.name=='Heading 1':
  num+=1;para.text=re.sub(r'^\\d+ ',str(num)+' ',para.text)
for para in doc.paragraphs:
 if 'CAN' in para.text and '150ms' in para.text:
  para.text='RES通信接口仅说明本模块自身的状态输出和故障处理；本报告不涉及外部控制器的集成或软件监督实现。'
# 调整叙述中的控制器名称为职责边界，只保留带下划线等现有信号名。
for el in doc._element.body.xpath('.//w:t'):
 if not el.text:continue
 el.text=re.sub(r'(?<![A-Za-z0-9_/])VCU(?![A-Za-z0-9_/])','外部控制器',el.text)
'''
t=t.replace(needle,post+'\n'+needle)
p.write_text(t,encoding='utf8')
print('Corrected ROTATE pins from newest netlist and excluded VCU source material')
