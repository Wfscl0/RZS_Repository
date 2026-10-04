from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml.ns import qn


SOURCE = Path(r"C:\Users\ABC\Downloads\附件1：个人测评报告（2025-2026学年）.docx")
OUTPUT = Path(r"C:\Users\ABC\Desktop\RZS_Repository\RES\附件1_个人测评报告_陈垚彬_已填写.docx")


def clear_paragraph(paragraph):
    """Remove runs and drawings while preserving the paragraph properties."""
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def add_formatted_text(paragraph, text, *, underline=True, bold=False):
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    run.font.name = "隶书"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "隶书")
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "隶书")
    run.font.size = Pt(12)
    run.font.underline = underline
    run.bold = bold
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 2.0


def remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)


doc = Document(SOURCE)

# Replace the identity paragraph while retaining the supplied template title and logo.
add_formatted_text(
    doc.paragraphs[3],
    "我是陈垚彬，学号为2024211992，合肥工业大学智能制造工程专业24-2班学生，现就个人2025-2026学年在政治素养、纪律意识、诚信记录、集体意识方面的情况总结如下：",
)

section_labels = {
    4: "政治素养方面：",
    8: "纪律意识方面：",
    11: "诚信记录方面：",
    16: "集体意识方面：",
}

# Keep the original template's section headings and underlined writing lines.
for index, text in section_labels.items():
    add_formatted_text(doc.paragraphs[index], text)

body_lines = {
    5: "本学年，我认真学习党的基本理论和路线方针政策，关心国家发展和社会热点，",
    6: "积极参加学校、学院和班级组织的学习教育活动，努力提高思想认识和政治素养。",
    7: "在日常学习生活中，我能够保持积极向上的态度，树立正确的价值观，端正成长方向。",
    9: "我能够自觉遵守国家法律法规和学校各项规章制度，遵守课堂纪律、考试纪律和宿舍管理要求，",
    10: "按时完成学习任务，合理安排时间，做到自律自觉、文明守纪，认真改进自身不足。",
    12: "我始终重视诚信品格养成，能够端正学习态度，独立完成作业和学习任务，",
    13: "遵守考试纪律，坚持实事求是，不弄虚作假，不抄袭，按要求如实反馈个人情况。",
    14: "在与老师和同学相处过程中，我能够做到诚恳待人、言行一致，认真履行应尽责任。",
    15: "今后我将继续把诚信要求落实到学习、生活和集体协作的各个方面，保持良好记录。",
    17: "我能够服从学校、学院和班级的工作安排，积极配合集体事务，按时完成分配的任务，",
    18: "与同学友好相处，主动沟通、互相帮助，努力维护班级团结和正常的学习生活秩序。",
    19: "在集体活动和团队协作中，我重视集体荣誉，能够顾全大局，认真承担自己的责任。",
    20: "今后我将进一步增强集体意识和服务意识，积极参与班级建设，在团队中不断提升自己。",
}

for index, text in body_lines.items():
    # Preserve the original underlined writing-line appearance by keeping a short
    # underlined tail after each completed line.
    add_formatted_text(doc.paragraphs[index], text + " " * 42)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
